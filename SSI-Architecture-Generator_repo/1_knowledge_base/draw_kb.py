"""Draw the knowledge base as it is filled (section 4 of the poster): where each sheet comes from and the relations between
FR, patterns and NFR, with the number of rows of each sheet. Reads SSI_Pattern_KnowledgeBase.xlsx only.
usage: python draw_kb.py  -> kb.png (Graphviz needed)"""
import os, shutil, subprocess
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
wb = openpyxl.load_workbook(os.path.join(HERE, "SSI_Pattern_KnowledgeBase.xlsx"), read_only=True)
n = {sh: sum(1 for r in wb[sh].iter_rows(values_only=True) if any(r)) - 1 for sh in wb.sheetnames}
n["NFR_proposed"] = sum(1 for r in wb["NFR"].iter_rows(min_row=2, values_only=True) if r[0] and "propos" in str(r[4]).lower())
cats = sorted({r[2] for r in list(wb["P"].iter_rows(values_only=True))[1:] if r[2]})

src = 'shape=note, style="filled", fillcolor="#f3f4f6", color="#6b7280", fontsize=16, margin="0.15,0.08"'
ent = 'shape=box, style="rounded,filled", fillcolor="#f7ebdd", color="#5b2a12", penwidth=2, fontsize=18, margin="0.2,0.1"'
L = ['digraph KB { rankdir=TB; nodesep=0.6; ranksep=0.55; splines=true; fontname=Helvetica; node [fontname=Helvetica]; '
     'edge [fontname=Helvetica, fontsize=16, color="#374151", fontcolor="#5b2a12", penwidth=1.6];',
     f'life [{src}, label="Lifecycles of keys, DIDs, VCs, VPs\\n[1], W3C: pre / postconditions"];',
     f'cat [{src}, label="SSI pattern catalogue [1]\\n35 patterns in prose, {len(cats)} categories\\ncontext, problem, solution, forces"];',
     f'cls [{src}, label="SSI properties [2]\\nclassification"];',
     f'FR [{ent}, label=<<B>FR</B>  {n["FR"]} functions>];',
     f'P [{ent}, label=<<B>P</B>  {n["P"]} patterns (35 + 5 variants)>];',
     f'NFR [{ent}, label=<<B>NFR</B>  {n["NFR"]} properties<BR/><FONT POINT-SIZE="18">{n["NFR"] - n["NFR_proposed"]} from [2] + {n["NFR_proposed"]} proposed</FONT>>];',
     'life -> FR [label=" one FR per transition", style=dashed]; cat -> P [label=" extract", style=dashed]; cls -> NFR [label=" extract", style=dashed];',
     f'FR -> P [label="realised_by ({n["realised_by"]})\\nmay_use ({n["may_use"]})"];',
     f'P -> NFR [label="affects ({n["affects"]})\\n++  +  −  −−"];',
     # loops under FR and P (an FR depends on an FR, a P requires a P), drawn through an invisible point below the node
     'dFR [shape=point, width=0.01, style=invis]; dP [shape=point, width=0.01, style=invis];',
     f'FR:sw -> dFR [arrowhead=none, tailport=sw]; dFR -> FR:se [xlabel="depends ({n["depends"]})  ", headport=se];',
     f'P:sw -> dP [arrowhead=none, tailport=sw]; dP -> P:se [xlabel="requires ({n["requires"]})  ", headport=se];',
     '{rank=same; life; cat; cls} {rank=same; FR; P; NFR}',
     '}']
dot = os.path.join(HERE, "kb.dot"); open(dot, "w", encoding="utf-8").write("\n".join(L))
if shutil.which("dot"):
    subprocess.run(["dot", "-Tpng", "-Gdpi=200", dot, "-o", os.path.join(HERE, "kb.png")])
    print("written kb.png:", {k: n[k] for k in ("FR", "NFR", "P", "realised_by", "may_use", "requires", "depends", "affects")})
